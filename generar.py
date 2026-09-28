"""Genera los requerimientos de pago (Word + PDF) y el Excel de envío por WhatsApp.

Uso:  python generar.py
Entrada: datos/cartera.csv  (y opcionalmente recursos/firma.png, recursos/sello.png)
Salida:  salida/word/*.docx, salida/pdf/*.pdf, salida/ENVIO_WHATSAPP.xlsx
"""
import csv
import html
import shutil
import unicodedata
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table

BASE = Path(__file__).parent
DATOS = BASE / "datos" / "cartera.csv"
RECURSOS = BASE / "recursos"
SALIDA = BASE / "salida"

# --- Datos del documento (editar aquí si cambian) -------------------------
FECHA_DOC = "28 de septiembre de 2026"
FECHA_LIMITE = "viernes 16 de octubre de 2026"
FIRMANTE = "ABOG. CRISTIAN DE JESÚS HERNÁNDEZ DÍAZ"
CARGO = "Apoderado Legal"
DIRECCION_BUFETE = ("Lomas del Guijarro, Zona Payaquí, atrás del Edificio NIVO, Casa 3640, "
                    "Tegucigalpa, Honduras.")
TEL_BUFETE = "9655-6803"
EMAIL_BUFETE = "amabledeje@yahoo.es"
FIRMA_MENSAJE = "Abg. Cristian Hernández, LEX-S Asesoría & Consultoría"

FINANCIERAS = {
    "CREDIMOVIL": {"razon": "CREDIMOVIL S.A. DE C.V.", "corto": "Credi Móvil", "sucursal": "Tegucigalpa"},
    "CREDIRAPID": {"razon": "CREDI RAPID S.A. DE C.V.", "corto": "Credi Rapid", "sucursal": "Tegucigalpa"},
    "PRESTAYA": {"razon": "PRESTA YA S.A. DE C.V.", "corto": "Presta Ya", "sucursal": "Tegucigalpa"},
    "PRESTAAUTO_CBA": {"razon": "PRESTA AUTO S.A. DE C.V.", "corto": "Presta Auto", "sucursal": "La Ceiba"},
}

FUNDAMENTO = (
    "Conforme a los artículos 1351, 1352, 1356 y 1361 del Código Civil de Honduras, las "
    "obligaciones contractuales deben cumplirse en los términos pactados. En caso de "
    "incumplimiento, la acreedora queda facultada para reclamar judicialmente el pago del "
    "capital adeudado, intereses, costas, gastos y demás cargos aplicables."
)


def lempiras(valor):
    return f"L. {float(valor):,.2f}"


def sin_acentos(texto):
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def apellidos(nombre):
    partes = nombre.split()
    return " ".join(p.capitalize() for p in partes[-2:])


def nombre_propio(nombre):
    return " ".join(p.capitalize() for p in nombre.split())


def texto_monto(r):
    tipo, monto = r["tipo_monto"], r["monto"]
    if tipo == "capital":
        return (f"usted mantiene un saldo de capital adeudado de {lempiras(monto)}, según los "
                "registros de la financiera, más los intereses corrientes y moratorios, cargos y "
                "gastos que correspondan conforme a la certificación de saldos")
    if tipo == "cancelacion":
        return (f"el saldo para cancelación total de su obligación asciende a {lempiras(monto)}, "
                "según los registros de la financiera, más los intereses moratorios, cargos y "
                "gastos que se continúen generando hasta su pago efectivo")
    if tipo == "mora":
        return (f"usted mantiene un saldo en mora (cuotas vencidas y no pagadas) de {lempiras(monto)}, "
                "más los intereses moratorios, cargos y gastos que correspondan, sin perjuicio del "
                "saldo total del contrato")
    return ("usted mantiene un saldo pendiente de pago cuyo monto total, con capital, intereses, "
            "cargos y gastos, constará en la certificación de saldos que emita la financiera")


def referencia(r, fin):
    partes = []
    if r["prestamo"]:
        ref = f"Contrato No. {r['prestamo']}"
        if r["perfil"]:
            ref += f" (perfil {r['perfil']})"
        partes.append(ref)
    partes.append(f"{fin['razon']}, Sucursal {fin['sucursal']}")
    veh = f"Vehículo {r['vehiculo'].title()}, " if r["vehiculo"] else "Vehículo "
    if r["placa"]:
        partes.append(f"{veh}Placa {r['placa']}")
    return " - ".join(partes) + "."


def cuerpo(r, fin):
    return (f"Por este medio le comunicamos que, ante el incumplimiento del contrato antes "
            f"referido, {fin['razon']} ha decidido reclamar judicialmente lo adeudado, siendo que "
            f"a la fecha {texto_monto(r)}.")


def ultimatum():
    return ("Previo a interponer la demanda judicial correspondiente, se le concede una última "
            "oportunidad para cancelar lo adeudado o formalizar un acuerdo de pago directo con la "
            f"financiera, a más tardar el {FECHA_LIMITE}. En caso de no hacerlo, el monto reclamado "
            "podrá verse incrementado dentro del proceso judicial por nuevos intereses, cargos, "
            "costas, gastos y demás conceptos que resulten aplicables conforme a derecho.")


def parrafo(doc, texto, negrita=False, tam=11, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, espacio=8):
    p = doc.add_paragraph()
    p.alignment = alineacion
    p.paragraph_format.space_after = Pt(espacio)
    run = p.add_run(texto)
    run.bold = negrita
    run.font.size = Pt(tam)
    return p


def crear_docx(r, fin, ruta):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(27.94), Cm(21.59)  # carta
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Times New Roman"
    estilo.font.size = Pt(11)

    membrete = RECURSOS / "membrete.png"
    if membrete.exists():
        doc.add_picture(str(membrete), width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        t = parrafo(doc, "LEX-S-", True, 20, WD_ALIGN_PARAGRAPH.CENTER, 0)
        t.runs[0].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        parrafo(doc, "ASESORÍA & CONSULTORÍA · DERECHO EMPRESARIAL Y POLÍTICAS PÚBLICAS",
                True, 9, WD_ALIGN_PARAGRAPH.CENTER, 14)

    parrafo(doc, "REQUERIMIENTO FORMAL DE PAGO", True, 14, WD_ALIGN_PARAGRAPH.CENTER, 2)
    parrafo(doc, f"Requerimiento Previo a la Vía Judicial - {fin['razon']}", False, 11,
            WD_ALIGN_PARAGRAPH.CENTER, 14)
    parrafo(doc, f"Tegucigalpa, M.D.C., {FECHA_DOC}.", alineacion=WD_ALIGN_PARAGRAPH.RIGHT)
    parrafo(doc, f"Señor(a): {r['cliente']}", True, espacio=2)
    if r["direccion"]:
        parrafo(doc, f"Dirección: {r['direccion']}", espacio=8)
    p = parrafo(doc, "")
    p.add_run("Referencia: ").bold = True
    p.add_run(referencia(r, fin))
    parrafo(doc, f"Estimado(a) señor(a) {apellidos(r['cliente'])}:")
    parrafo(doc, cuerpo(r, fin))
    parrafo(doc, FUNDAMENTO)
    parrafo(doc, ultimatum())
    parrafo(doc, "Atentamente,", espacio=4)

    firma, sello = RECURSOS / "firma.png", RECURSOS / "sello.png"
    if firma.exists() or sello.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if firma.exists():
            p.add_run().add_picture(str(firma), height=Cm(2.2))
        if sello.exists():
            p.add_run("   ")
            p.add_run().add_picture(str(sello), height=Cm(3.0))
    else:
        for _ in range(3):
            doc.add_paragraph()
        parrafo(doc, "_________________________________", alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio=0)

    parrafo(doc, FIRMANTE, True, alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio=0)
    parrafo(doc, CARGO, alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio=0)
    parrafo(doc, fin["razon"], True, alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio=14)
    pie = parrafo(doc, f"{DIRECCION_BUFETE}  Cel.: {TEL_BUFETE}  ·  Email: {EMAIL_BUFETE}",
                  tam=8, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    pie.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.save(ruta)


def estilo_pdf(nombre, tam=11, negrita=False, alineacion=TA_JUSTIFY, despues=8, color="#000000"):
    return ParagraphStyle(nombre, fontName="Times-Bold" if negrita else "Times-Roman",
                          fontSize=tam, leading=tam * 1.3, alignment=alineacion,
                          spaceAfter=despues, textColor=HexColor(color))


def crear_pdf(r, fin, ruta):
    normal = estilo_pdf("n")
    centro = estilo_pdf("c", alineacion=TA_CENTER, despues=0)
    centro_b = estilo_pdf("cb", negrita=True, alineacion=TA_CENTER, despues=0)
    e = []
    membrete = RECURSOS / "membrete.png"
    if membrete.exists():
        e.append(Image(str(membrete), width=16.5 * cm, height=16.5 * cm * 0.18, kind="proportional"))
    else:
        e.append(Paragraph("LEX-S-", estilo_pdf("m", 20, True, TA_CENTER, 2, "#1F3A5F")))
        e.append(Paragraph("ASESORÍA &amp; CONSULTORÍA · DERECHO EMPRESARIAL Y POLÍTICAS PÚBLICAS",
                           estilo_pdf("m2", 9, True, TA_CENTER, 14)))
    e.append(Paragraph("REQUERIMIENTO FORMAL DE PAGO", estilo_pdf("t", 14, True, TA_CENTER, 2)))
    e.append(Paragraph(f"Requerimiento Previo a la Vía Judicial - {fin['razon']}",
                       estilo_pdf("st", 11, False, TA_CENTER, 14)))
    e.append(Paragraph(f"Tegucigalpa, M.D.C., {FECHA_DOC}.", estilo_pdf("f", alineacion=TA_RIGHT)))
    e.append(Paragraph(f"<b>Señor(a): {html.escape(r['cliente'])}</b>", estilo_pdf("s", despues=2)))
    if r["direccion"]:
        e.append(Paragraph(f"Dirección: {html.escape(r['direccion'])}", normal))
    e.append(Paragraph(f"<b>Referencia:</b> {html.escape(referencia(r, fin))}", normal))
    e.append(Paragraph(f"Estimado(a) señor(a) {apellidos(r['cliente'])}:", normal))
    for texto in (cuerpo(r, fin), FUNDAMENTO, ultimatum()):
        e.append(Paragraph(texto, normal))
    e.append(Paragraph("Atentamente,", normal))

    firma, sello = RECURSOS / "firma.png", RECURSOS / "sello.png"
    imagenes = []
    if firma.exists():
        imagenes.append(Image(str(firma), width=5 * cm, height=2.2 * cm, kind="proportional"))
    if sello.exists():
        imagenes.append(Image(str(sello), width=3.2 * cm, height=3.2 * cm, kind="proportional"))
    if imagenes:
        e.append(Table([imagenes]))
    else:
        e.append(Spacer(1, 1.6 * cm))
        e.append(Paragraph("_________________________________", centro))
    e.append(Paragraph(FIRMANTE, centro_b))
    e.append(Paragraph(CARGO, centro))
    e.append(Paragraph(fin["razon"], centro_b))
    e.append(Spacer(1, 0.6 * cm))
    e.append(Paragraph(f"{DIRECCION_BUFETE}  Cel.: {TEL_BUFETE}  ·  Email: {EMAIL_BUFETE}",
                       estilo_pdf("pie", 8, False, TA_CENTER, 0, "#555555")))
    SimpleDocTemplate(str(ruta), pagesize=letter, leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm, pageCompression=1,
                      title=f"Requerimiento de pago - {r['cliente']}", author="LEX-S").build(e)


def telefono_wa(numero):
    digitos = "".join(c for c in numero if c.isdigit())
    return f"504{digitos}" if len(digitos) == 8 else ""


def mensaje_wa(r, fin):
    nombre = nombre_propio(r["cliente"])
    if r["tipo_monto"] == "sin_monto":
        saldo = "el saldo pendiente de su contrato"
    elif r["tipo_monto"] == "mora":
        saldo = f"el saldo en mora de {lempiras(r['monto'])} de su contrato"
    else:
        saldo = f"el saldo de {lempiras(r['monto'])} de su contrato"
    placa = f" (vehículo placa {r['placa']})" if r["placa"] else ""
    return (f"Buenos días, {nombre}. Le hago llegar el REQUERIMIENTO FORMAL DE PAGO previo a la vía "
            f"judicial de {fin['razon']}{placa}, por {saldo}. Tiene plazo hasta el {FECHA_LIMITE} "
            f"para pagar o formalizar un arreglo de pago con la financiera; de lo contrario se "
            f"presentará la demanda judicial. Le adjunto el documento. Quedo atento a su "
            f"comunicación. {FIRMA_MENSAJE}. Cel. {TEL_BUFETE}.")


def crear_excel(filas, ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "ENVIAR HOY"
    encabezados = ["No.", "QUÉ HACER", "Financiera", "Cliente", "Celular", "ABRIR WHATSAPP",
                   "Celular 2", "ABRIR WHATSAPP 2", "Saldo reclamado", "Placa", "Préstamo",
                   "PDF (arrastrar al chat)", "Mensaje (ya va escrito en el link)",
                   "Observaciones", "¿Enviado?", "Fecha/hora envío"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A5F")
        celda.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    colores = {"ENVIAR": "C6EFCE", "CONFIRMAR": "FFEB9C", "NO ENVIAR": "FFC7CE", "PEDIR CELULAR": "FFEB9C"}
    enlace = Font(color="0563C1", underline="single", bold=True)
    for i, (r, fin, pdf) in enumerate(filas, start=2):
        msg = mensaje_wa(r, fin)
        tel1, tel2 = telefono_wa(r["celular"]), telefono_wa(r["celular2"])
        que_hacer = {
            "ENVIAR": "ENVIAR",
            "CONFIRMAR": "CONFIRMAR CON AMABLE antes de enviar",
            "NO ENVIAR": "NO ENVIAR",
            "PEDIR CELULAR": "PEDIR CELULAR a la financiera",
        }[r["accion"]]
        saldo = "Según certificación" if r["tipo_monto"] == "sin_monto" else float(r["monto"])
        ws.append([i - 1, que_hacer, fin["corto"], nombre_propio(r["cliente"]), r["celular"],
                   "Abrir chat" if tel1 else "", r["celular2"], "Abrir chat" if tel2 else "",
                   saldo, r["placa"], r["prestamo"], pdf, msg, r["observacion"], "No", ""])
        if tel1:
            ws.cell(i, 6).hyperlink = f"https://wa.me/{tel1}?text={quote(msg)}"
            ws.cell(i, 6).font = enlace
        if tel2:
            ws.cell(i, 8).hyperlink = f"https://wa.me/{tel2}?text={quote(msg)}"
            ws.cell(i, 8).font = enlace
        ws.cell(i, 12).hyperlink = f"pdf/{pdf}"
        ws.cell(i, 9).number_format = '"L. "#,##0.00'
        ws.cell(i, 2).fill = PatternFill("solid", fgColor=colores[r["accion"]])
        ws.cell(i, 2).font = Font(bold=True)
        for c in range(1, len(encabezados) + 1):
            ws.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")

    total = len(filas) + 1
    validacion = DataValidation(type="list", formula1='"No,Sí,No tiene WhatsApp,Entregado en físico"')
    ws.add_data_validation(validacion)
    validacion.add(f"O2:O{total}")
    anchos = [5, 22, 12, 30, 12, 13, 12, 13, 15, 10, 9, 45, 60, 40, 12, 16]
    for n, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(1, n).column_letter].width = ancho
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:P{total}"

    inst = wb.create_sheet("CÓMO ENVIAR")
    pasos = [
        "CÓMO ENVIAR (2 minutos por cliente)",
        "1. Abra WhatsApp Web (web.whatsapp.com) en Chrome con el celular del bufete.",
        "2. En la hoja ENVIAR HOY, haga clic en 'Abrir chat': se abre el chat con el mensaje YA ESCRITO.",
        "3. Arrastre al chat el PDF de la columna 'PDF' (carpeta pdf, junto a este Excel) y presione Enviar.",
        "4. Si WhatsApp dice que el número no está en WhatsApp, pruebe 'ABRIR WHATSAPP 2' o márquelo 'No tiene WhatsApp'.",
        "5. Ponga 'Sí' en ¿Enviado? y la fecha/hora; tome captura de pantalla (prueba del requerimiento).",
        "",
        "HORARIO: domingo 9:00 a.m. a 1:00 p.m.; lunes a sábado 8:00 a.m. a 8:00 p.m.; feriados NO.",
        "Filas AMARILLAS: confirmar con Amable antes de enviar. Filas ROJAS: NO enviar.",
        "Guarde las capturas: algunos juzgados piden acreditar el requerimiento (incluso con acta notarial).",
    ]
    for p in pasos:
        inst.append([p])
    inst["A1"].font = Font(bold=True, size=14)
    inst.column_dimensions["A"].width = 120

    resumen = wb.create_sheet("RESUMEN")
    resumen.append(["Financiera", "Requerimientos", "Enviar", "Confirmar / pedir dato", "No enviar", "Total reclamado (L.)"])
    for c in resumen[1]:
        c.font = Font(bold=True)
    for clave, fin in FINANCIERAS.items():
        grupo = [r for r, f, _ in filas if f is fin]
        resumen.append([
            fin["corto"], len(grupo),
            sum(r["accion"] == "ENVIAR" for r in grupo),
            sum(r["accion"] in ("CONFIRMAR", "PEDIR CELULAR") for r in grupo),
            sum(r["accion"] == "NO ENVIAR" for r in grupo),
            sum(float(r["monto"]) for r in grupo if r["monto"]),
        ])
        resumen.cell(resumen.max_row, 6).number_format = '#,##0.00'
    for col, ancho in zip("ABCDEF", [14, 15, 8, 22, 10, 20]):
        resumen.column_dimensions[col].width = ancho
    wb.save(ruta)


def unir_pdf(archivos, ruta):
    try:
        import pymupdf
    except ImportError:
        print("Instale pymupdf para generar TODOS_PARA_IMPRIMIR.pdf (pip install pymupdf)")
        return
    unido = pymupdf.open()
    for archivo in archivos:
        unido.insert_pdf(pymupdf.open(archivo))
    unido.save(ruta, garbage=4, deflate=True)


def main():
    if SALIDA.exists():
        shutil.rmtree(SALIDA)
    word, pdf = SALIDA / "word", SALIDA / "pdf"
    word.mkdir(parents=True)
    pdf.mkdir()

    with open(DATOS, encoding="utf-8-sig") as f:
        registros = list(csv.DictReader(f))

    filas = []
    for n, r in enumerate(registros, start=1):
        fin = FINANCIERAS[r["financiera"]]
        base = sin_acentos(f"{n:02d}_Requerimiento_{r['cliente'].replace(' ', '_')}_{fin['corto'].replace(' ', '')}")
        crear_docx(r, fin, word / f"{base}.docx")
        crear_pdf(r, fin, pdf / f"{base}.pdf")
        filas.append((r, fin, f"{base}.pdf"))

    crear_excel(filas, SALIDA / "ENVIO_WHATSAPP.xlsx")
    unir_pdf([pdf / p for _, _, p in filas], SALIDA / "TODOS_PARA_IMPRIMIR.pdf")
    faltan = [p for _, _, p in filas if not (pdf / p).exists()]
    print(f"{len(filas)} requerimientos generados en {SALIDA}")
    if faltan:
        print("PDF que no se generaron:", faltan)


if __name__ == "__main__":
    main()
